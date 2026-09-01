#!/usr/bin/env python3
"""Build public/sku_model_list.xlsx from named tables in SharePoint."""

from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import quote

GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
GRAPH_SCOPE = "https://graph.microsoft.com/.default"
REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG_PATH = REPO_ROOT / "config" / "sharepoint-mapping.json"
OUTPUT_HEADERS = ("sku", "model", "raw_model")


class SyncError(RuntimeError):
    """Raised when source data cannot safely produce the matcher workbook."""


@dataclass(frozen=True)
class TableConfig:
    name: str
    sku_column: str
    value_column: str


@dataclass(frozen=True)
class SyncConfig:
    site_hostname: str
    site_path: str
    file_path: str
    raw_model_table: TableConfig
    model_table: TableConfig
    output_path: Path


@dataclass(frozen=True)
class TableData:
    headers: tuple[str, ...]
    rows: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class BuildStats:
    output_rows: int
    matched_raw_models: int
    missing_raw_models: int
    param_only_skus: int
    skipped_blank_models: int


def main() -> int:
    args = parse_args()
    config = load_config(args.config)
    output_path = Path(args.output).resolve() if args.output else config.output_path

    if args.input_workbook:
        print(f"Reading local workbook fixture: {args.input_workbook}")
        workbook_bytes = Path(args.input_workbook).read_bytes()
    else:
        print(f"SharePoint site: https://{config.site_hostname}{config.site_path}")
        print(f"SharePoint file: {config.file_path}")
        workbook_bytes = download_sharepoint_workbook(config)

    table_names = (config.raw_model_table.name, config.model_table.name)
    print(f"Named tables: {', '.join(table_names)}")
    tables = load_named_tables(workbook_bytes, table_names)
    rows, stats = build_mapping_rows(
        tables[config.raw_model_table.name],
        tables[config.model_table.name],
        config.raw_model_table,
        config.model_table,
    )

    existing_rows = read_existing_mapping(output_path)
    if existing_rows == rows and not args.force_write:
        print(
            f"No data changes. Kept {output_path} unchanged "
            f"({stats.output_rows} rows, {stats.missing_raw_models} blank raw_model values)."
        )
        return 0

    write_mapping_workbook(output_path, rows)
    print(
        f"Updated {output_path}: rows={stats.output_rows}, "
        f"matched_raw_model={stats.matched_raw_models}, "
        f"blank_raw_model={stats.missing_raw_models}, "
        f"param_only_not_exported={stats.param_only_skus}, "
        f"blank_model_not_exported={stats.skipped_blank_models}"
    )
    return 0


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--config",
        default=str(DEFAULT_CONFIG_PATH),
        help="Committed JSON config containing non-secret SharePoint routing.",
    )
    parser.add_argument(
        "--input-workbook",
        help="Optional local Bava_data.xlsx fixture; bypasses SharePoint authentication.",
    )
    parser.add_argument("--output", help="Optional output path override.")
    parser.add_argument(
        "--force-write",
        action="store_true",
        help="Rewrite the output even when its row data is unchanged.",
    )
    return parser.parse_args()


def load_config(path_value: str | Path) -> SyncConfig:
    config_path = Path(path_value).resolve()
    try:
        payload = json.loads(config_path.read_text(encoding="utf-8"))
        sharepoint = payload["sharepoint"]
        tables = payload["tables"]
        raw_model = tables["raw_model"]
        model = tables["model"]
        output_value = payload["output_path"]
    except (OSError, json.JSONDecodeError, KeyError, TypeError) as exc:
        raise SyncError(f"Invalid sync config {config_path}: {exc}") from exc

    output_path = Path(str(output_value))
    if not output_path.is_absolute():
        output_path = REPO_ROOT / output_path

    return SyncConfig(
        site_hostname=required_config_value(sharepoint, "site_hostname"),
        site_path=normalize_site_path(required_config_value(sharepoint, "site_path")),
        file_path=required_config_value(sharepoint, "file_path"),
        raw_model_table=TableConfig(
            name=required_config_value(raw_model, "name"),
            sku_column=required_config_value(raw_model, "sku_column"),
            value_column=required_config_value(raw_model, "value_column"),
        ),
        model_table=TableConfig(
            name=required_config_value(model, "name"),
            sku_column=required_config_value(model, "sku_column"),
            value_column=required_config_value(model, "value_column"),
        ),
        output_path=output_path.resolve(),
    )


def required_config_value(container: dict[str, Any], key: str) -> str:
    value = str(container.get(key, "")).strip()
    if not value:
        raise SyncError(f"Sync config value '{key}' must not be blank.")
    return value


def normalize_site_path(value: str) -> str:
    path = "/" + value.strip().strip("/")
    return path.rstrip("/")


def required_env(name: str) -> str:
    value = os.getenv(name, "").strip()
    if not value:
        raise SyncError(f"Missing required environment variable: {name}")
    return value


def download_sharepoint_workbook(config: SyncConfig) -> bytes:
    tenant_id = required_env("TENANT_ID")
    client_id = required_env("CLIENT_ID")
    client_secret = required_env("CLIENT_SECRET")
    token = get_graph_access_token(tenant_id, client_id, client_secret)

    site_payload = graph_request_json(
        "GET",
        f"{GRAPH_BASE_URL}/sites/{config.site_hostname}:{config.site_path}?$select=id",
        token,
    )
    site_id = str(site_payload.get("id", "")).strip()
    if not site_id:
        raise SyncError("Microsoft Graph did not return a SharePoint site id.")

    drive_payload = graph_request_json(
        "GET",
        f"{GRAPH_BASE_URL}/sites/{site_id}/drive?$select=id,name",
        token,
    )
    drive_id = str(drive_payload.get("id", "")).strip()
    if not drive_id:
        raise SyncError("Microsoft Graph did not return a document-library drive id.")

    drive_path = normalize_drive_path(config.file_path)
    encoded_path = quote(drive_path, safe="/")
    return graph_request_bytes(
        "GET",
        f"{GRAPH_BASE_URL}/drives/{drive_id}/root:/{encoded_path}:/content",
        token,
        timeout=120,
    )


def get_graph_access_token(tenant_id: str, client_id: str, client_secret: str) -> str:
    import requests

    try:
        response = requests.post(
            f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
            headers={"Content-Type": "application/x-www-form-urlencoded"},
            data={
                "client_id": client_id,
                "client_secret": client_secret,
                "grant_type": "client_credentials",
                "scope": GRAPH_SCOPE,
            },
            timeout=60,
        )
    except requests.RequestException as exc:
        raise SyncError(f"Microsoft Graph token request failed: {exc}") from exc
    if response.status_code != 200:
        raise SyncError(
            f"Microsoft Graph token request failed with HTTP {response.status_code}: "
            f"{response.text[:500]}"
        )
    token = str(response.json().get("access_token", "")).strip()
    if not token:
        raise SyncError("Microsoft Graph token response did not include access_token.")
    return token


def normalize_drive_path(value: str) -> str:
    path = value.strip().replace("\\", "/").lstrip("/")
    while "//" in path:
        path = path.replace("//", "/")
    lowered = path.casefold()
    for prefix in ("shared documents/", "documents/"):
        if lowered.startswith(prefix):
            return path[len(prefix) :]
    return path


def graph_request_json(method: str, url: str, token: str) -> dict[str, Any]:
    response = graph_request(method, url, token, timeout=60)
    try:
        return response.json()
    except ValueError as exc:
        raise SyncError(f"Microsoft Graph returned invalid JSON for {url}") from exc


def graph_request_bytes(
    method: str,
    url: str,
    token: str,
    *,
    timeout: int,
) -> bytes:
    return graph_request(method, url, token, timeout=timeout).content


def graph_request(
    method: str,
    url: str,
    token: str,
    *,
    timeout: int,
) -> Any:
    import requests

    backoff_seconds = 1.0
    for attempt in range(1, 7):
        try:
            response = requests.request(
                method,
                url,
                headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
                timeout=timeout,
            )
        except requests.RequestException as exc:
            raise SyncError(f"Microsoft Graph request failed: {exc}") from exc
        if response.status_code < 400:
            return response
        if response.status_code not in {429, 500, 502, 503, 504}:
            raise SyncError(
                f"Microsoft Graph request failed with HTTP {response.status_code}: "
                f"{response.text[:500]}"
            )
        if attempt == 6:
            break
        retry_after = response.headers.get("Retry-After", "")
        try:
            wait_seconds = max(float(retry_after), 0.0) if retry_after else backoff_seconds
        except ValueError:
            wait_seconds = backoff_seconds
        print(
            f"Microsoft Graph returned HTTP {response.status_code}; "
            f"retrying in {wait_seconds:.1f}s ({attempt}/6)."
        )
        time.sleep(wait_seconds)
        backoff_seconds = min(backoff_seconds * 2, 20.0)
    raise SyncError("Microsoft Graph request failed after six attempts.")


def load_named_tables(
    workbook_bytes: bytes,
    requested_names: tuple[str, ...],
) -> dict[str, TableData]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=True, read_only=False)
    requested_by_key = {canonical(name): name for name in requested_names}
    found: dict[str, TableData] = {}
    known_names: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            for table_name in worksheet.tables.keys():
                known_names.append(table_name)
                requested_name = requested_by_key.get(canonical(table_name))
                if not requested_name:
                    continue
                table = worksheet.tables[table_name]
                cells = worksheet[table.ref]
                matrix = [[cell.value for cell in row] for row in cells]
                if not matrix:
                    raise SyncError(f"Named table '{requested_name}' is empty.")
                headers = tuple(clean_string(value) for value in matrix[0])
                rows: list[dict[str, Any]] = []
                for values in matrix[1:]:
                    row = {
                        headers[index]: value
                        for index, value in enumerate(values)
                        if index < len(headers) and headers[index]
                    }
                    if any(clean_string(value) for value in row.values()):
                        rows.append(row)
                found[requested_name] = TableData(headers=headers, rows=tuple(rows))
    finally:
        workbook.close()

    missing = [name for name in requested_names if name not in found]
    if missing:
        raise SyncError(
            f"Named table(s) not found: {', '.join(missing)}. "
            f"Workbook tables: {', '.join(sorted(known_names)) or '[none]'}"
        )
    return found


def build_mapping_rows(
    raw_model_data: TableData,
    model_data: TableData,
    raw_model_config: TableConfig,
    model_config: TableConfig,
) -> tuple[list[dict[str, str]], BuildStats]:
    raw_model_by_sku, _, _ = index_table(
        raw_model_data,
        raw_model_config,
        require_value=False,
        skip_blank_value=False,
    )
    model_by_sku, model_order, skipped_blank_models = index_table(
        model_data,
        model_config,
        require_value=True,
        skip_blank_value=True,
    )

    if not model_order:
        raise SyncError(f"Named table '{model_config.name}' has no valid product rows.")

    rows = [
        {
            "sku": sku,
            "model": model_by_sku[sku],
            "raw_model": raw_model_by_sku.get(sku, ""),
        }
        for sku in model_order
    ]
    validate_output_rows(rows)

    matched = sum(1 for row in rows if row["raw_model"])
    param_only = len(set(raw_model_by_sku) - set(model_by_sku))
    stats = BuildStats(
        output_rows=len(rows),
        matched_raw_models=matched,
        missing_raw_models=len(rows) - matched,
        param_only_skus=param_only,
        skipped_blank_models=skipped_blank_models,
    )
    return rows, stats


def index_table(
    table: TableData,
    config: TableConfig,
    *,
    require_value: bool,
    skip_blank_value: bool,
) -> tuple[dict[str, str], list[str], int]:
    sku_header = resolve_header(table.headers, config.sku_column, config.name)
    value_header = resolve_header(table.headers, config.value_column, config.name)
    indexed: dict[str, str] = {}
    order: list[str] = []
    skipped_blank_values = 0

    for row_number, row in enumerate(table.rows, start=2):
        sku = clean_sku(row.get(sku_header))
        value = clean_string(row.get(value_header))
        if not sku:
            raise SyncError(
                f"Table '{config.name}' has a non-empty row without SKU at table row {row_number}."
            )
        if not value:
            if skip_blank_value:
                skipped_blank_values += 1
                continue
            if require_value:
                raise SyncError(
                    f"Table '{config.name}' has blank '{config.value_column}' for SKU {sku}."
                )
        if sku in indexed:
            raise SyncError(f"Table '{config.name}' contains duplicate SKU {sku}.")
        indexed[sku] = value
        order.append(sku)

    return indexed, order, skipped_blank_values


def resolve_header(headers: tuple[str, ...], wanted: str, table_name: str) -> str:
    by_key = {canonical(header): header for header in headers if header}
    resolved = by_key.get(canonical(wanted))
    if not resolved:
        raise SyncError(
            f"Table '{table_name}' is missing column '{wanted}'. "
            f"Available columns: {', '.join(headers)}"
        )
    return resolved


def canonical(value: Any) -> str:
    return re.sub(r"\s+", "", clean_string(value).casefold())


def clean_string(value: Any) -> str:
    return "" if value is None else str(value).strip()


def clean_sku(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def validate_output_rows(rows: list[dict[str, str]]) -> None:
    seen: set[str] = set()
    for index, row in enumerate(rows, start=2):
        if tuple(row.keys()) != OUTPUT_HEADERS:
            raise SyncError(f"Output row {index} does not match required column order.")
        if not row["sku"] or not row["model"]:
            raise SyncError(f"Output row {index} has a blank sku or model.")
        if row["sku"] in seen:
            raise SyncError(f"Output contains duplicate SKU {row['sku']}.")
        seen.add(row["sku"])


def read_existing_mapping(path: Path) -> list[dict[str, str]] | None:
    from openpyxl import load_workbook

    if not path.exists():
        return None
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            return None
        headers = tuple(clean_string(value).casefold() for value in header_values[:3])
        if headers != OUTPUT_HEADERS:
            return None
        existing: list[dict[str, str]] = []
        for values in rows:
            sku = clean_sku(values[0] if len(values) > 0 else None)
            model = clean_string(values[1] if len(values) > 1 else None)
            raw_model = clean_string(values[2] if len(values) > 2 else None)
            if not sku and not model and not raw_model:
                continue
            existing.append({"sku": sku, "model": model, "raw_model": raw_model})
        return existing
    finally:
        workbook.close()


def write_mapping_workbook(path: Path, rows: list[dict[str, str]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    worksheet.sheet_view.showGridLines = False
    worksheet.append(list(OUTPUT_HEADERS))

    for row in rows:
        worksheet.append([row["sku"], row["model"], row["raw_model"]])

    thin = Side(style="thin", color="FFB7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_font = Font(name="Arial", size=12)
    header_font = Font(name="Arial", size=12, bold=True)

    for cell in worksheet[1]:
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.font = body_font
            cell.border = border
        row[0].number_format = "@"
        row[0].alignment = Alignment(horizontal="center", vertical="center")

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 110
    worksheet.column_dimensions["C"].width = 45
    worksheet.auto_filter.ref = f"A1:C{len(rows) + 1}"

    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}.",
            suffix=path.suffix,
            dir=path.parent,
            delete=False,
        ) as temp_file:
            temp_path = Path(temp_file.name)
        workbook.save(temp_path)
        os.replace(temp_path, path)
    finally:
        workbook.close()
        if temp_path and temp_path.exists():
            temp_path.unlink()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (SyncError, OSError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2) from exc
