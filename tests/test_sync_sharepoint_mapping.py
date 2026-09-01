from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from scripts.sync_sharepoint_mapping import (
    SyncError,
    TableConfig,
    TableData,
    build_mapping_rows,
    clean_sku,
    load_named_tables,
    normalize_drive_path,
    read_existing_mapping,
    write_mapping_workbook,
)


RAW_CONFIG = TableConfig("_param_table", "SKU", "Model")
MODEL_CONFIG = TableConfig("_rus_kaz_table", "SKU", "Описание")


class MappingBuildTests(unittest.TestCase):
    def test_rus_kaz_is_base_and_param_is_left_joined(self) -> None:
        raw_data = TableData(
            headers=("SKU", "Model"),
            rows=(
                {"SKU": 100, "Model": "Raw 100"},
                {"SKU": 300, "Model": "Raw 300"},
            ),
        )
        model_data = TableData(
            headers=("SKU", "Описание"),
            rows=(
                {"SKU": 100, "Описание": "Description 100"},
                {"SKU": 200, "Описание": "Description 200"},
                {"SKU": 400, "Описание": ""},
            ),
        )

        rows, stats = build_mapping_rows(raw_data, model_data, RAW_CONFIG, MODEL_CONFIG)

        self.assertEqual(
            rows,
            [
                {"sku": "100", "model": "Description 100", "raw_model": "Raw 100"},
                {"sku": "200", "model": "Description 200", "raw_model": ""},
            ],
        )
        self.assertEqual(stats.output_rows, 2)
        self.assertEqual(stats.matched_raw_models, 1)
        self.assertEqual(stats.missing_raw_models, 1)
        self.assertEqual(stats.param_only_skus, 1)
        self.assertEqual(stats.skipped_blank_models, 1)

    def test_duplicate_sku_fails_instead_of_overwriting(self) -> None:
        raw_data = TableData(
            headers=("SKU", "Model"),
            rows=({"SKU": "100", "Model": "Raw 100"},),
        )
        model_data = TableData(
            headers=("SKU", "Описание"),
            rows=(
                {"SKU": "100", "Описание": "Description A"},
                {"SKU": "100", "Описание": "Description B"},
            ),
        )

        with self.assertRaisesRegex(SyncError, "duplicate SKU 100"):
            build_mapping_rows(raw_data, model_data, RAW_CONFIG, MODEL_CONFIG)

    def test_missing_required_column_fails(self) -> None:
        raw_data = TableData(
            headers=("SKU", "Wrong"),
            rows=({"SKU": "100", "Wrong": "Raw 100"},),
        )
        model_data = TableData(
            headers=("SKU", "Описание"),
            rows=({"SKU": "100", "Описание": "Description 100"},),
        )

        with self.assertRaisesRegex(SyncError, "missing column 'Model'"):
            build_mapping_rows(raw_data, model_data, RAW_CONFIG, MODEL_CONFIG)

    def test_sharepoint_library_prefix_is_not_part_of_drive_path(self) -> None:
        self.assertEqual(
            normalize_drive_path(
                "/Shared Documents/General/_system_files/Bava_data.xlsx"
            ),
            "General/_system_files/Bava_data.xlsx",
        )

    def test_numeric_sku_is_normalized_without_decimal_suffix(self) -> None:
        self.assertEqual(clean_sku(4513812.0), "4513812")
        self.assertEqual(clean_sku("4513812.0"), "4513812")

    def test_named_tables_are_extracted_and_output_round_trips(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.worksheet.table import Table

        with TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "Bava_data.xlsx"
            output_path = Path(temp_dir) / "sku_model_list.xlsx"

            workbook = Workbook()
            param_sheet = workbook.active
            param_sheet.title = "Parameters"
            param_sheet.append(["SKU", "Model"])
            param_sheet.append([4513812, "TE-HD 18 Li-Solo"])
            param_sheet.add_table(Table(displayName="_param_table", ref="A1:B2"))

            rus_sheet = workbook.create_sheet("Descriptions")
            rus_sheet.append(["SKU", "Описание"])
            rus_sheet.append(
                [4513812, "Аккум. Перфоратор Einhell TE-HD 18 Li-Solo (4513812)"]
            )
            rus_sheet.add_table(Table(displayName="_rus_kaz_table", ref="A1:B2"))
            workbook.save(source_path)
            workbook.close()

            tables = load_named_tables(
                source_path.read_bytes(),
                (RAW_CONFIG.name, MODEL_CONFIG.name),
            )
            rows, _ = build_mapping_rows(
                tables[RAW_CONFIG.name],
                tables[MODEL_CONFIG.name],
                RAW_CONFIG,
                MODEL_CONFIG,
            )
            write_mapping_workbook(output_path, rows)

            self.assertEqual(read_existing_mapping(output_path), rows)
            output_workbook = load_workbook(output_path, data_only=True, read_only=False)
            try:
                output_sheet = output_workbook.worksheets[0]
                self.assertEqual(
                    [cell.value for cell in output_sheet[1]],
                    ["sku", "model", "raw_model"],
                )
                self.assertEqual(output_sheet.auto_filter.ref, "A1:C2")
                self.assertFalse(output_sheet.sheet_view.showGridLines)
            finally:
                output_workbook.close()


if __name__ == "__main__":
    unittest.main()
